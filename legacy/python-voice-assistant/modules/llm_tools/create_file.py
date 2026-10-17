import google.generativeai as genai
from config.settings import GEMINI_API_KEY, CURRENT_OS ,DEFAULT_SAVE_PATH
import os
import re
import subprocess
import shlex
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from reportlab.lib.enums import TA_CENTER
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Pt

def _ensure_genai():
    if not GEMINI_API_KEY:
        raise RuntimeError("Gemini API key missing. Configure it in the installer or run the app's configuration.")
    genai.configure(api_key=GEMINI_API_KEY)

def open_file(filepath):
    """Open a file in its default application based on the current OS."""
    try:
        if CURRENT_OS.lower().startswith("win"):  # Windows
            os.startfile(filepath)
        elif CURRENT_OS.lower().startswith("darwin"):  # macOS
            subprocess.run(["open", filepath])
        elif CURRENT_OS.lower().startswith("linux"):  # Linux
            subprocess.run(["xdg-open", filepath])
        else:
            return "Unsupported OS for auto-opening file."
    except Exception as e:
        return " Could not open file automatically"

def file_generator(topic, filetype="pdf", save_path=None):
    _ensure_genai()
    model = genai.GenerativeModel("gemini-2.0-flash")
    response = model.generate_content(f"You are an AI content generator. Your task is to create detailed, structured, and professional content for a given topic. The generated text will be directly saved into a file (txt, pdf, docx, markdown, etc.), so follow these rules:1. Output ONLY the raw content of the file. Do not include explanations, notes, code fences, or extra text outside the content. 2. Organize the content with clear headings, sections, and paragraphs. 3. Write in a professional, informative, and detailed manner. 4. Ensure formatting is clean and suitable for saving in different file types (plain text, PDF, DOCX, or Markdown). 5. If lists, tables, or examples are needed, include them in a neat format.Topic: {topic},File Type: {filetype}"
)
    generated_text = response.text.strip()

    if save_path:
        os.makedirs(os.path.dirname(save_path), exist_ok=True)
        full_path = save_path
    else:
        os.makedirs(DEFAULT_SAVE_PATH, exist_ok=True)
        full_path = os.path.join(DEFAULT_SAVE_PATH, f"{topic.replace(' ','_')}.txt")

    with open(full_path, "w", encoding="utf-8") as f:
        f.write(generated_text)

    
    if filetype=="pdf":
        txt_to_pdf(full_path, os.path.join(DEFAULT_SAVE_PATH, f"{topic.replace(' ','_')}.pdf"))
    elif filetype=="docx":
        txt_to_docx(full_path, os.path.join(DEFAULT_SAVE_PATH, f"{topic.replace(' ','_')}.docx"))
    elif filetype=="pptx":
        txt_to_pptx(full_path, os.path.join(DEFAULT_SAVE_PATH, f"{topic.replace(' ','_')}.pptx"))
    








# ---------------- TXT → PDF ---------------- #
def txt_to_pdf(txt_file, pdf_file, title=" "):
    with open(txt_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = SimpleDocTemplate(pdf_file, pagesize=A4,
                            rightMargin=50, leftMargin=50,
                            topMargin=50, bottomMargin=50)

    styles = getSampleStyleSheet()
    styles["Heading1"].alignment = TA_CENTER
    story = [Paragraph(title, styles["Heading1"]), Spacer(1, 20)]

    list_buffer = []
    for line in lines:
        line = line.strip()
        if not line:
            continue

        # Headings
        if line.startswith("### "):
            story.append(Paragraph(line[4:], styles["Heading3"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], styles["Heading1"]))
        elif line.startswith("- "):
            list_buffer.append(ListItem(Paragraph(line[2:], styles["Normal"])))
        else:
            if list_buffer:
                story.append(ListFlowable(list_buffer, bulletType="bullet"))
                list_buffer = []
            line = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", line)
            line = re.sub(r"\*(.*?)\*", r"<i>\1</i>", line)
            story.append(Paragraph(line, styles["Normal"]))
            story.append(Spacer(1, 10))

    if list_buffer:
        story.append(ListFlowable(list_buffer, bulletType="bullet"))

    doc.build(story)
    open_file(pdf_file)
    return f" PDF saved: {pdf_file}"


# ---------------- TXT → DOCX ---------------- #
def txt_to_docx(txt_file, docx_file, title=" "):
    with open(txt_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    doc.add_heading(title, level=1)

    for line in lines:
        line = line.strip()
        if not line:
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="ListBullet")
        else:
            p = doc.add_paragraph()
            tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*)", line)
            for token in tokens:
                if not token:
                    continue
                run = p.add_run()
                if token.startswith("**") and token.endswith("**"):
                    run.text = token[2:-2]
                    run.bold = True
                elif token.startswith("*") and token.endswith("*"):
                    run.text = token[1:-1]
                    run.italic = True
                else:
                    run.text = token
                run.font.size = Pt(12)

    doc.save(docx_file)
    open_file(docx_file)
    return f" DOCX saved: {docx_file}"


# ---------------- TXT → PPTX ---------------- #
def txt_to_pptx(txt_file, pptx_file):
    with open(txt_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    prs = Presentation()
    MAX_PARAGRAPHS = 12  # Auto-split threshold

    current_slide = None
    text_frame = None

    def new_content_slide(title_text=""):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title_text
        return slide, slide.shapes.placeholders[1].text_frame

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # -------- TITLE SLIDE -------- #
        if line.startswith("# "):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = line[2:]
            current_slide, text_frame = None, None
            continue

        # -------- NEW SECTION / NEW SLIDE -------- #
        if line.startswith("## "):
            current_slide, text_frame = new_content_slide(line[3:])
            continue

        # If no slide exists yet, create one
        if not current_slide:
            current_slide, text_frame = new_content_slide("")

        # -------- AUTO-SPLIT SLIDE -------- #
        if len(text_frame.paragraphs) >= MAX_PARAGRAPHS:
            # Create a new continuation slide
            current_slide, text_frame = new_content_slide(current_slide.shapes.title.text + " (cont.)")

        # -------- SUBHEADING -------- #
        if line.startswith("### "):
            p = text_frame.add_paragraph()
            p.text = line[4:]
            p.font.bold = True
            p.font.size = Pt(22)
            continue

        # -------- BULLETS -------- #
        if line.startswith("- "):
            p = text_frame.add_paragraph()
            p.text = line[2:]
            p.level = 0
            continue

        # -------- NORMAL TEXT -------- #
        clean_line = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
        clean_line = re.sub(r"\*(.*?)\*", r"\1", clean_line)

        p = text_frame.add_paragraph()
        p.text = clean_line
        p.font.size = Pt(18)

    prs.save(pptx_file)
    open_file(pptx_file)
    return f"PPTX saved: {pptx_file}"
